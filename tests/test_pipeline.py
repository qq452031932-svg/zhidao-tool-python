from pathlib import Path
import json

import pytest

from lxds import pipeline
from lxds.excelio import read_resource_table, write_resource_table


def test_parse_simple_yaml(tmp_path):
    path = tmp_path / "pipeline.yaml"
    path.write_text('category: 小说\ncount: 2\ntransfer:\n  mode: external # C# 处理\n  external_inbox: "D:/exchange/in"\n  external_outbox: "D:/exchange/out"\n', encoding="utf-8")
    loaded = pipeline.load_pipeline_config(str(path))
    assert loaded["transfer"]["mode"] == "external"
    assert loaded["count"] == 2


def test_transfer_preserves_only_share_success(monkeypatch, tmp_path):
    source = tmp_path / "source.xlsx"
    target = tmp_path / "transferred.xlsx"
    write_resource_table(str(source), [{"name": "书甲", "link": "https://pan.baidu.com/s/1abc?pwd=abcd"}])

    class FakePan:
        def __init__(self, cookie):
            assert cookie == "secret"
        def transfer_link(self, url, target_path, do_share):
            assert do_share is True
            return {"status": "转存成功", "share_url": "https://pan.baidu.com/s/1new?pwd=1234", "share_pwd": "1234"}

    monkeypatch.setattr(pipeline, "PanClient", FakePan)
    result = pipeline.transfer_file(str(source), str(target), "/我的资源", True, "secret")
    rows, skipped = read_resource_table(str(target))
    assert result["count"] == 1
    assert skipped == []
    assert rows[0]["link"] == "https://pan.baidu.com/s/1new?pwd=1234"


def test_template_file_writes_qid_workbook(monkeypatch, tmp_path):
    source = tmp_path / "transferred.xlsx"
    target = tmp_path / "qa.xlsx"
    write_resource_table(str(source), [{"name": "书甲", "link": "https://pan.baidu.com/s/1abc?pwd=abcd"}])
    monkeypatch.setattr(pipeline, "ai_generate_batch", lambda key, rows: [
        {"title": "求书甲完整网盘链接下载", "answer": "<p>回答</p>"}])
    result = pipeline.template_file(str(source), str(target), "test-key")
    from openpyxl import load_workbook
    sheet = load_workbook(target, read_only=True).active
    assert result["count"] == 1
    assert [cell.value for cell in sheet[1]] == ["qid", "问题标题", "回答内容"]


def test_external_transfer_copies_and_reads_returned_workbook(monkeypatch, tmp_path):
    source = tmp_path / "collected.xlsx"
    output = tmp_path / "transferred.xlsx"
    inbox, outbox = tmp_path / "in", tmp_path / "out"
    write_resource_table(str(source), [{"name": "书甲", "link": "https://pan.baidu.com/s/1old?pwd=abcd"}])

    def produce(path, timeout, poll, after_ns=0, result_glob=None):
        write_resource_table(str(path), [{"name": "书甲", "link": "https://pan.baidu.com/s/1new?pwd=1234"}])
        return path
    monkeypatch.setattr(pipeline, "_wait_external", produce)
    result = pipeline.external_transfer(str(source), str(output), str(inbox), str(outbox), 1)
    assert (inbox / "collected.xlsx").exists()
    assert result["count"] == 1
    assert read_resource_table(str(output))[0][0]["link"].endswith("?pwd=1234")


def test_run_writes_all_intermediates_with_mocks(monkeypatch, tmp_path):
    def fake_mine(category, count, out, **kwargs):
        Path(out).write_text("小说书甲\n", encoding="utf-8")
        return {"count": 1}
    def fake_collect(mode, keywords, category, count, out, **kwargs):
        write_resource_table(out, [{"name": "书甲", "link": "https://pan.baidu.com/s/1old?pwd=abcd"}])
        return {"count": 1}
    def fake_transfer(source, out, target, share, cookie):
        write_resource_table(out, [{"name": "书甲", "link": "https://pan.baidu.com/s/1new?pwd=1234"}])
        return {"count": 1}
    def fake_template(source, out, key):
        Path(out).write_bytes(b"xlsx")
        return {"count": 1}
    def fake_publish(source, account):
        assert Path(source).read_bytes() == b"xlsx"
        return {"task_id": "ok"}
    monkeypatch.setattr(pipeline, "mine_file", fake_mine)
    monkeypatch.setattr(pipeline, "collect_file", fake_collect)
    monkeypatch.setattr(pipeline, "transfer_file", fake_transfer)
    monkeypatch.setattr(pipeline, "template_file", fake_template)
    monkeypatch.setattr(pipeline, "publish_file", fake_publish)
    monkeypatch.setattr(pipeline, "_credentials", lambda: {"pan_cookie": "p", "deepseek_key": "k", "zhidao_accounts": []})
    result = pipeline.run_pipeline({"category": "小说", "count": 1, "workdir": str(tmp_path),
                                    "transfer": {"mode": "local", "target": "/我的资源"}, "publish": {"account": "账号甲"}})
    assert result["published"]["task_id"] == "ok"
    assert all((tmp_path / name).exists() for name in ("keywords.txt", "collected.xlsx", "transferred.xlsx", "qa.xlsx"))


def test_all_cli_subcommands_parse():
    cases = [
        ["mine", "--category", "小说", "--count", "2", "--out", "keywords.txt"],
        ["collect", "--mode", "precise", "--keywords", "keywords.txt", "--out", "collected.xlsx"],
        ["transfer", "--in", "collected.xlsx", "--target", "/我的资源", "--share", "--out", "transferred.xlsx"],
        ["template", "--in", "transferred.xlsx", "--out", "qa.xlsx"],
        ["publish", "--in", "qa.xlsx", "--account", "账号甲"],
        ["answer", "--question", "哪里有", "--style", "first_person"],
        ["run", "--config", "pipeline.yaml"],
    ]
    for argv in cases:
        assert pipeline.build_parser().parse_args(argv).command == argv[0]


def test_external_wait_ignores_stale_file(tmp_path):
    path = tmp_path / "collected.xlsx"
    path.write_bytes(b"old")
    with pytest.raises(TimeoutError):
        pipeline._wait_external(path, 0.02, 0.005, after_ns=path.stat().st_mtime_ns)


def test_run_uses_real_excel_transfer_and_template_steps(monkeypatch, tmp_path):
    from lxds import miner
    from openpyxl import load_workbook

    monkeypatch.setitem(miner.SOURCES, "小说", (("fake", lambda *a: ["书甲"], False),))
    monkeypatch.setattr(miner, "_search_supplement", lambda *a: [])

    def fake_collect(mode, keywords, category, count, out, **kwargs):
        assert Path(keywords).read_text(encoding="utf-8").strip() == "小说书甲"
        write_resource_table(out, [{"name": "书甲", "link": "https://pan.baidu.com/s/1old?pwd=abcd"}])
        return {"count": 1}

    class FakePan:
        def __init__(self, cookie):
            assert cookie == "p"
        def transfer_link(self, url, target, share):
            assert share is True
            return {"status": "转存成功", "share_url": "https://pan.baidu.com/s/1new?pwd=1234", "share_pwd": "1234"}

    monkeypatch.setattr(pipeline, "collect_file", fake_collect)
    monkeypatch.setattr(pipeline, "PanClient", FakePan)
    monkeypatch.setattr(pipeline, "ai_generate_batch", lambda key, rows: [
        {"title": "求书甲完整网盘链接下载", "answer": "<p>回答</p>"}])
    monkeypatch.setattr(pipeline, "publish_file", lambda path, account: {"task_id": account})
    monkeypatch.setattr(pipeline, "_credentials", lambda: {"pan_cookie": "p", "deepseek_key": "k", "zhidao_accounts": []})

    result = pipeline.run_pipeline({"category": "小说", "count": 1, "workdir": str(tmp_path),
                                    "transfer": {"mode": "local", "target": "/资源"}, "publish": {"account": "甲"}})
    assert result["published"] == {"task_id": "甲"}
    assert read_resource_table(str(tmp_path / "transferred.xlsx"))[0][0]["link"].endswith("?pwd=1234")
    assert load_workbook(tmp_path / "qa.xlsx", read_only=True).active.max_row == 2
