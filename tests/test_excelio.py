"""M02 三列契约与百度知道上传模板的验收测试。"""

from io import BytesIO

import openpyxl
import pytest

from lxds.excelio import (
    build_publish_workbook_bytes,
    read_resource_table,
    write_resource_table,
)


LINK = "https://pan.baidu.com/s/1AbCd?pwd=abcd"


def _save_sheet(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


@pytest.mark.parametrize(
    ("head", "expected_name"),
    [
        (["资源名称", "资源链接", "提取码"], "数学试卷"),
        (["源标题", "分享链接", "提取码", "状态", "账号"], "数学试卷"),
        (["资源名称", "分享链接", "提取码", "状态", "账号", "目标目录", "备注", "错误"], "数学试卷"),
    ],
)
def test_reads_named_headers_and_ignores_extra_columns(tmp_path, head, expected_name):
    path = tmp_path / "items.xlsx"
    _save_sheet(path, [head, ["数学试卷", LINK, "abcd", "转存成功", "甲"]])
    rows, skipped = read_resource_table(str(path))
    assert rows == [{"name": expected_name, "link": LINK, "index": 2}]
    assert skipped == []


def test_reads_headerless_abc_and_fills_missing_name(tmp_path):
    path = tmp_path / "items.xlsx"
    _save_sheet(path, [[None, LINK, "abcd"]])
    rows, skipped = read_resource_table(str(path))
    assert rows == [{"name": "资源1", "link": LINK, "index": 1}]
    assert skipped == []


def test_detects_header_on_third_row(tmp_path):
    path = tmp_path / "items.xlsx"
    _save_sheet(path, [["说明"], ["导出时间"], ["文件名字", "分享链接", "提取码"], ["古诗合集", LINK, "abcd"]])
    rows, skipped = read_resource_table(str(path))
    assert rows == [{"name": "古诗合集", "link": LINK, "index": 4}]
    assert skipped == []


def test_skips_blank_and_non_baidu_links_with_reasons(tmp_path):
    path = tmp_path / "items.xlsx"
    _save_sheet(path, [["资源名称", "资源链接"], ["甲", None], ["乙", "https://example.com/a"]])
    rows, skipped = read_resource_table(str(path))
    assert rows == []
    assert len(skipped) == 2
    assert "链接为空" in skipped[0]
    assert "不是百度网盘链接" in skipped[1]


def test_write_resource_table_has_fixed_headers_and_widths(tmp_path):
    path = tmp_path / "out.xlsx"
    write_resource_table(str(path), [{"name": "古诗合集", "link": LINK, "pwd": "abcd"}])
    sheet = openpyxl.load_workbook(path).active
    assert [sheet.cell(1, column).value for column in (1, 2, 3)] == ["资源名称", "资源链接", "提取码"]
    assert [sheet.cell(2, column).value for column in (1, 2, 3)] == ["古诗合集", LINK, "abcd"]
    assert [sheet.column_dimensions[column].width for column in "ABC"] == [40, 60, 16]


def test_publish_workbook_has_empty_qid_and_exact_layout():
    content = build_publish_workbook_bytes([{"title": "求古诗合集完整网盘链接下载", "answer": "<p>内容</p>"}])
    sheet = openpyxl.load_workbook(BytesIO(content)).active
    assert [sheet.cell(1, column).value for column in (1, 2, 3)] == ["qid", "问题标题", "回答内容"]
    assert [sheet.cell(2, column).value for column in (1, 2, 3)] == [None, "求古诗合集完整网盘链接下载", "<p>内容</p>"]
    assert [sheet.column_dimensions[column].width for column in "ABC"] == [8, 50, 80]
    assert sheet.cell(1, 1).font.bold
    assert sheet.cell(1, 1).alignment.horizontal == "center"


def test_csv_round_trip_and_xls_error(tmp_path):
    path = tmp_path / "out.csv"
    write_resource_table(str(path), [{"name": "古诗合集", "link": LINK, "pwd": "abcd"}])
    rows, skipped = read_resource_table(str(path))
    assert rows == [{"name": "古诗合集", "link": LINK, "index": 2}]
    assert skipped == []
    with pytest.raises(ValueError, match="请用 Excel 另存为 .xlsx 后上传"):
        read_resource_table(str(tmp_path / "old.xls"))


def test_separate_extraction_code_is_added_to_link(tmp_path):
    path = tmp_path / "items.xlsx"
    _save_sheet(path, [["资源名称", "资源链接", "提取码"], ["古诗合集", "https://pan.baidu.com/s/1AbCd", "abcd"]])
    rows, skipped = read_resource_table(str(path))
    assert rows[0]["link"] == LINK
    assert skipped == []


def test_legacy_replacement_result_prefers_new_share_link(tmp_path):
    path = tmp_path / "legacy.xlsx"
    old_link = "https://pan.baidu.com/s/1Old?pwd=old1"
    new_link = "https://pan.baidu.com/s/1New?pwd=new1"
    _save_sheet(path, [
        ["第 1 列", "第 2 列", "转存结果_解析后标题", "转存结果_原分享链接", "转存结果_新分享链接"],
        ["旧标题", old_link, "解析后标题", old_link, new_link],
        ["失败标题", old_link, "失败标题", old_link, None],
    ])
    rows, skipped = read_resource_table(str(path))
    assert rows == [{"name": "解析后标题", "link": new_link, "index": 2}]
    assert len(skipped) == 1 and "链接为空" in skipped[0]
