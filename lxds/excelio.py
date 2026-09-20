"""资源表和百度知道批量发布模板的 Excel/CSV 契约。"""

from __future__ import annotations

import csv
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


_NAME_HEADERS = ("名称", "文件名", "名字", "资源名", "源标题", "解析后标题")
_LINK_HEADERS = ("链接", "网盘")
_PAN_LINK = re.compile(r"^https?://pan\.baidu\.com/s/[0-9A-Za-z_-]+(?:\?.*)?$", re.I)
_PWD = re.compile(r"(?:^|[?&])pwd=([0-9A-Za-z]{4})(?:&|$)", re.I)
_RESOURCE_HEADERS = ("资源名称", "资源链接", "提取码")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _value(row: list[Any], index: int | None) -> str:
    return _text(row[index]) if index is not None and index < len(row) else ""


def _detect_columns(rows: list[list[Any]]) -> tuple[int, int, int | None, int]:
    """在前五行找表头；找不到时按无表头 A/B/C 读。"""
    for offset, row in enumerate(rows[:5]):
        cells = [_text(value) for value in row]
        name = next((i for i, value in enumerate(cells) if "解析后标题" in value), None)
        if name is None:
            name = next((i for i, value in enumerate(cells) if any(word in value for word in _NAME_HEADERS)), None)
        # 旧程序结果表同时有原链接与新链接，必须用转存后的新链接。
        link = next((i for i, value in enumerate(cells) if "新分享链接" in value), None)
        if link is None:
            link = next((i for i, value in enumerate(cells) if any(word in value for word in _LINK_HEADERS)), None)
        if name is None:
            continue
        if link is None:
            # 某些来源的表头不含“链接”，用该列中的网盘地址定位。
            for candidate in range(max((len(item) for item in rows[offset + 1:offset + 6]), default=0)):
                if candidate != name and any("pan.baidu.com/s/" in _value(item, candidate) for item in rows[offset + 1:offset + 6]):
                    link = candidate
                    break
        if link is None or link == name:
            continue
        pwd = next((i for i, value in enumerate(cells) if "提取码" in value or "访问码" in value), None)
        return name, link, pwd, offset + 1
    return 0, 1, 2, 0


def _read_rows(path: Path) -> list[list[Any]]:
    extension = path.suffix.lower()
    if extension == ".xls":
        raise ValueError("请用 Excel 另存为 .xlsx 后上传")
    if extension == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            return [list(row) for row in csv.reader(file)]
    if extension == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    raise ValueError(f"不支持的文件格式：{extension or '无扩展名'}")


def _link_with_pwd(link: str, pwd: str) -> str:
    if pwd and re.fullmatch(r"[0-9A-Za-z]{4}", pwd) and not _PWD.search(link):
        return f"{link}{'&' if '?' in link else '?'}pwd={pwd}"
    return link


def read_resource_table(path: str) -> tuple[list[dict], list[str]]:
    """读资源表，返回 `({name, link, index} 列表, 跳过原因列表)`。"""
    source_rows = _read_rows(Path(path))
    name_col, link_col, pwd_col, start = _detect_columns(source_rows)
    rows: list[dict] = []
    skipped: list[str] = []
    for number, source in enumerate(source_rows[start:], start=start + 1):
        if not any(_text(cell) for cell in source):
            continue
        link = _value(source, link_col)
        if not link:
            skipped.append(f"第{number}行：链接为空")
            continue
        if not _PAN_LINK.fullmatch(link):
            skipped.append(f"第{number}行：不是百度网盘链接")
            continue
        link = _link_with_pwd(link, _value(source, pwd_col))
        name = _value(source, name_col) or f"资源{number}"
        rows.append({"name": name, "link": link, "index": number})
    return rows, skipped


def _get(item: dict, *names: str) -> str:
    for name in names:
        value = _text(item.get(name))
        if value:
            return value
    return ""


def write_resource_table(path: str, items: list[dict]) -> None:
    """严格输出三列资源表；支持 xlsx 和 csv。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    rows = []
    for item in items:
        name = _get(item, "name", "资源名称")
        link = _get(item, "link", "资源链接", "分享链接")
        match = _PWD.search(link)
        pwd = _get(item, "pwd", "提取码") or (match.group(1) if match else "")
        rows.append((name, _link_with_pwd(link, pwd), pwd))
    extension = target.suffix.lower()
    if extension == ".csv":
        with target.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(_RESOURCE_HEADERS)
            writer.writerows(rows)
        return
    if extension != ".xlsx":
        raise ValueError("资源表仅支持 .xlsx 或 .csv")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(_RESOURCE_HEADERS)
    for row in rows:
        sheet.append(row)
    for column, width in {"A": 40, "B": 60, "C": 16}.items():
        sheet.column_dimensions[column].width = width
    workbook.save(target)


def build_publish_workbook_bytes(rows: list[dict]) -> bytes:
    """内存生成百度知道要求的 qid/标题/回答三列 xlsx。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("qid", "问题标题", "回答内容"))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append((None, _get(row, "title", "问题标题"), _get(row, "answer", "回答内容")))
    for column, width in {"A": 8, "B": 50, "C": 80}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
